import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from bulk_export import save_bulk_json, write_bulk_excel


def sample_batch():
    settings = {
        "starting_cash": 100, "target_profit": 101.25, "max_loss": 17.5,
        "max_rounds": 500, "mode": "custom", "betting_sequence": [1, 2],
        "max_picks": 2, "difficulty": "low", "grinding_enabled": False,
        "grinding_range": 2, "p_random": False,
        "win_multipliers": {"low": {"1": 1.25, "2": 1.5}},
        "mine_config": {"low": 1}, "custom_mode": [[1, 2, "low"], [2, 1, "low"]],
        "betting_modes": {"normal": [1]}, "bet_values": [1, 2], "label": "Custom",
    }
    seed = 9223372036854770000
    runs = [
        {"index": 1, "seed": seed, "status": "target", "rounds": 8,
         "final_cash": 210, "profit": 110, "max_drawdown": 4, "peak_cash": 210,
         "threshold_drawdowns": [0, 4]},
        {"index": 2, "seed": seed + 1, "status": "max_loss", "rounds": 3,
         "final_cash": 80, "profit": -20, "max_drawdown": 20, "peak_cash": 100,
         "threshold_drawdowns": [None, None]},
        {"index": 3, "seed": seed + 2, "status": "interrupted", "rounds": 1,
         "final_cash": 99, "profit": -1, "max_drawdown": 1, "peak_cash": 100,
         "threshold_drawdowns": [None, None]},
    ]
    summary = {
        "completed_runs": 2, "interrupted_runs": 1, "target_rate": 50,
        "positive_rate": 50, "avg_profit": 45, "median_profit": 45,
        "min_profit": -20, "max_profit": 110, "p05_profit": -13.5,
        "p95_profit": 103.5, "avg_rounds": 5.5, "max_drawdown": 20,
        "target_ci_low": 9.453, "target_ci_high": 90.547,
        "positive_ci_low": 9.453, "positive_ci_high": 90.547,
        "stop_counts": {"target": 1, "max_loss": 1, "max_rounds": 0,
                        "insufficient_funds": 0, "interrupted": 1},
    }
    config = {"id": 1, "label": "Custom", "settings": settings,
              "profit_thresholds": [1, 101.25], "loss_thresholds": [5, 17.5],
              "success_counts": [[1, 1], [1, 1]], "runs": runs, "summary": summary}
    return {"schema_version": 1, "created_at": "2026-09-19T12:00:00+00:00", "seed": seed,
            "requested_runs": 10, "status": "interrupted", "elapsed_seconds": 1.25,
            "configurations": [config]}


class BulkExportTest(unittest.TestCase):
    def test_workbook_preserves_metrics_partial_runs_and_color_grid(self):
        batch = sample_batch()
        second = copy.deepcopy(batch["configurations"][0])
        second.update(id=2, label="=literal label")
        second["runs"] = []
        second["success_counts"] = [[0, 0], [0, 0]]
        for key in second["summary"]:
            second["summary"][key] = None
        second["summary"].update(completed_runs=0, interrupted_runs=0, stop_counts={})
        batch["configurations"].append(second)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "nested" / "results.xlsx"
            write_bulk_excel(path, batch)
            workbook = load_workbook(path)
            self.addCleanup(workbook.close)
            self.assertEqual(workbook.sheetnames, [
                "Confronto laptop", "Griglia Win Rate", "Impostazioni", "Sommario", "Confronto",
                "Dettaglio simulazioni", "Distribuzione profitti",
            ])
            grid = workbook["Griglia Win Rate"]
            self.assertEqual(grid["C5"].value, 101.25)
            self.assertEqual(grid["A7"].value, 17.5)
            self.assertEqual(grid["B6"].value, 50)  # 1 / 2, not 1 / 3.
            self.assertIn("%", grid["B6"].number_format)
            self.assertIsNone(grid["B11"].value)
            self.assertEqual(len(grid.conditional_formatting), 2)
            self.assertEqual(grid.freeze_panes, "B6")
            comparison = workbook["Confronto"]
            headers = {cell.value: cell.column for cell in comparison[1]}
            self.assertEqual(comparison.cell(2, headers["Profitto medio"]).value, 45)
            self.assertEqual(comparison.cell(2, headers["Target: IC 95% inferiore (%)"]).value, 9.453)
            self.assertEqual(comparison.cell(2, headers["Perdita massima raggiunta"]).value, 1)
            self.assertEqual(comparison.cell(2, headers["Interrotta dall'utente"]).value, 1)
            self.assertIsNone(comparison.cell(3, headers["Target raggiunto (%)"]).value)
            self.assertEqual(comparison["B3"].value, "=literal label")
            self.assertEqual(comparison["B3"].data_type, "s")
            self.assertIn("ConfrontoConfigurazioni", comparison.tables)
            self.assertTrue(comparison.tables["ConfrontoConfigurazioni"].autoFilter.ref)
            details = workbook["Dettaglio simulazioni"]
            self.assertEqual(details.max_row, 4)
            self.assertNotIn("Seed", str([cell.value for cell in details[1]]))
            self.assertEqual(details["D4"].value, "Interrotta dall'utente")
            self.assertEqual(details["G4"].value, -1)
            metadata = {row[0]: row[1] for row in workbook["Impostazioni"].iter_rows(values_only=True) if row[0]}
            self.assertEqual(json.loads(metadata["Moltiplicatori per difficoltà / picks"])["low"]["2"], 1.5)
            histogram = workbook["Distribuzione profitti"]
            self.assertEqual(len(histogram._charts), 1)
            frequencies = [row[3] for row in histogram.iter_rows(values_only=True)
                           if isinstance(row[3], int)]
            self.assertEqual(sum(frequencies), 2)

    def test_empty_and_interrupted_only_batches_have_no_fake_rates_or_histograms(self):
        for has_partial_run in (False, True):
            with self.subTest(has_partial_run=has_partial_run), tempfile.TemporaryDirectory() as directory:
                batch = sample_batch()
                config = batch["configurations"][0]
                config["runs"] = config["runs"][2:] if has_partial_run else []
                config["success_counts"] = [[0, 0], [0, 0]]
                summary = {key: None for key in config["summary"]}
                summary.update(completed_runs=0, interrupted_runs=int(has_partial_run),
                               stop_counts={"interrupted": int(has_partial_run)})
                config["summary"] = summary
                path = Path(directory) / "empty.xlsx"
                write_bulk_excel(path, batch)
                workbook = load_workbook(path)
                try:
                    self.assertIsNone(workbook["Griglia Win Rate"]["B6"].value)
                    self.assertEqual(len(workbook["Distribuzione profitti"]._charts), 0)
                    self.assertEqual(workbook["Dettaglio simulazioni"].max_row, 1 + int(has_partial_run))
                finally:
                    workbook.close()

    def test_json_roundtrip_and_failed_replacement_preserve_existing_snapshot(self):
        batch = sample_batch()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "nested" / "snapshot.json"
            save_bulk_json(path, batch)
            self.assertEqual(json.loads(path.read_text(encoding="utf-8")), batch)
            original = path.read_bytes()
            batch["status"] = "completed"
            with patch("bulk_export.os.replace", side_effect=PermissionError("locked")):
                with self.assertRaises(PermissionError):
                    save_bulk_json(path, batch)
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(list(path.parent.iterdir()), [path])


if __name__ == "__main__":
    unittest.main()
