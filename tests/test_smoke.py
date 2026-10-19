r"""Run with: .venv\Scripts\python.exe -m unittest discover -s tests"""

import asyncio
import json
from pathlib import Path
import tempfile
import tkinter as tk
from tkinter import ttk
import unittest
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from game_engine import GameEngine
from gratta_e_vinci_gui import GrattaEVinciGUI
from settings_manager import SettingsManager


class SmokeTest(unittest.TestCase):
    def test_gui_simulation_and_monitor_lifecycle(self):
        root = tk.Tk()
        root.withdraw()
        self.addCleanup(lambda: root.destroy() if root.tk.call('info', 'commands', '.') else None)
        callback_errors = []
        root.report_callback_exception = lambda *error: callback_errors.append(error)
        defaults = SettingsManager().merge_with_defaults({})
        with patch('pyautogui.click', side_effect=AssertionError('Unexpected real click')), \
             patch('pyautogui.moveTo', side_effect=AssertionError('Unexpected mouse move')), \
             patch('pyautogui.position', return_value=(123, 456)) as position, \
             patch.object(SettingsManager, 'load_settings', return_value=defaults):
            app = GrattaEVinciGUI(root)
            self.assertEqual(app.get_win_multiplier('medium', 4), 3.2)
            self.assertEqual(app.get_win_multiplier('medium', 6), 6.8)
            self.assertEqual(app.get_win_multiplier('medium', 7), 10.0)
            self.assertEqual(app.get_win_multiplier('medium', 8), 15.0)
            self.assertEqual(int(app.max_picks_spinbox.cget('to')), 8)
            for panel in app.bulk_laptops:
                self.assertIn('8', panel.game_controls[1].cget('values'))
            previous_mode, previous_custom = app.mode_var.get(), app.custom_mode
            app.mode_var.set('custom')
            app.custom_mode = [{'b': .1, 'p': 8, 'd': 'medium'}]
            self.assertTrue(app.game_engine.validate_settings())
            self.assertEqual(app.game_engine._build_bulk_snapshot()['win_multipliers']['medium'][8], 15.0)
            round_config = app.game_engine._get_round_config_for_strategy()
            self.assertEqual((round_config['max_picks'], round_config['mine_count'], round_config['multiplier']),
                             (8, 6, 15.0))
            app.custom_mode = [{'b': .1, 'p': 8, 'd': 'low'}]
            with self.assertRaises(ValueError):
                app.game_engine.validate_settings()
            app.mode_var.set(previous_mode)
            app.custom_mode = previous_custom
            # Loading older settings keeps the new confirmed medium quotes.
            old_payload = app._collect_settings_payload()
            for picks in (5, 6, 7, 8):
                del old_payload['bulk_game_rules']['win_multipliers']['medium'][picks]
            app._apply_settings_payload(old_payload)
            self.assertEqual(app.bulk_win_multiplier_vars['medium'][6].get(), 6.8)
            self.assertEqual(app.bulk_win_multiplier_vars['medium'][7].get(), 10.0)
            self.assertEqual(app.bulk_win_multiplier_vars['medium'][8].get(), 15.0)
            notebook = next(widget for widget in root.winfo_children() if isinstance(widget, ttk.Notebook))
            self.assertEqual([notebook.tab(tab, 'text') for tab in notebook.tabs()],
                             ['Settings', 'Bulk Test', 'Coordinates', 'Betting Modes', 'Inizializzazione',
                              'Game Control', 'Statistics', 'Pause'])
            settings = notebook.nametowidget(notebook.tabs()[1])
            for widget in (app.bulk_start_button, app.bulk_stop_button, app.bulk_log_text,
                           app.bulk_n_runs_entry, app.bulk_export_button):
                self.assertTrue(str(widget).startswith(str(settings) + '.'))
            root.after(250, root.quit)
            root.mainloop()
            self.assertGreaterEqual(position.call_count, 2)
            self.assertIn('123, 456', app.coord_label.cget('text'))

            app.max_rounds_var.set(3)
            app.game_running = True
            app.running_event.set()
            app.game_engine.initialize_game_variables()
            app.game_engine.run_test_mode()
            root.update()
            self.assertEqual(app.rounds, 3)
            self.assertFalse(app.game_running)
            self.assertIn('Reached maximum rounds', app.progress_label.cget('text'))

            app.starting_cash_var.set(100)
            app.target_win_var.set(101)
            app.max_loss_var.set(5)
            app.max_rounds_var.set(20)
            app.max_picks_var.set(1)
            app.bulk_betting_modes['normal'] = [1.0]
            for variable in app.bulk_mine_config_vars.values():
                variable.set(0)
            results = app.game_engine.run_bulk_test(3, lambda *_: None, lambda _: None,
                                                    lambda: False, lambda: False)
            config = results['configurations'][0]
            self.assertEqual(config['summary']['completed_runs'], 3)
            self.assertEqual(config['success_counts'], [[3]])
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / 'results.xlsx'
                app._write_excel(path, results)
                workbook = load_workbook(path)
                self.assertEqual(workbook.sheetnames, ['Confronto laptop', 'Griglia Win Rate', 'Impostazioni', 'Sommario',
                                                     'Confronto', 'Dettaglio simulazioni', 'Distribuzione profitti'])
                self.assertEqual(workbook['Griglia Win Rate']['B6'].value, 100)
                workbook.close()

                # Exercise three independent panels, worker, plain JSON and Excel.
                app.bulk_output_dir_var.set(directory)
                app.bulk_n_runs_var.set(2)
                self.assertEqual(len(app.bulk_laptops), 3)
                for index, panel in enumerate(app.bulk_laptops, 1):
                    panel.values['mode'].set('normal')
                    panel.values['difficulty'].set('low')
                    panel.values['max_picks'].set('1')
                    panel.values['starting_cash'].set('100')
                    panel.values['target_profit'].set(str(index))
                    panel.values['max_loss'].set('10')
                    panel.values['max_rounds'].set('20')
                    panel.values['grinding_enabled'].set(False)
                    panel.mode_sequences['normal'] = [1.0]
                app.bulk_win_multiplier_vars['low'][1].set(2.0)
                # Editing laptop 1 must not change laptop 2's own sequence.
                app.bulk_laptops[0].mode_sequences['normal'].append(2.0)
                self.assertEqual(app.bulk_laptops[1].mode_sequences['normal'], [1.0])
                self.assertEqual([p.snapshot()['label'] for p in app.bulk_laptops], ['Laptop 1', 'Laptop 2', 'Laptop 3'])
                self.assertFalse(hasattr(app, 'bulk_seed_var'))
                self.assertFalse(hasattr(app, 'bulk_compare_var'))
                payload = app._collect_settings_payload()
                self.assertEqual(len(payload['bulk_laptops']), 3)
                app.bulk_laptops[0].values['target_profit'].set('99')
                app._apply_settings_payload(payload)
                self.assertEqual(float(app.bulk_laptops[0].values['target_profit'].get()), 1)
                timeout = root.after(10000, root.quit)
                def await_completion():
                    if app.bulk_test_running:
                        root.after(20, await_completion)
                    else:
                        root.quit()
                with patch.object(app, '_show_bulk_grid_popup') as show_results, \
                     patch('gratta_e_vinci_gui.messagebox.showerror') as show_error:
                    app.start_bulk_test()
                    root.after(20, await_completion)
                    root.mainloop()
                    root.after_cancel(timeout)
                    self.assertFalse(app.bulk_test_running)
                    show_error.assert_not_called()
                    show_results.assert_called_once()
                self.assertEqual(len(app._bulk_results['configurations']), 3)
                self.assertEqual([c['summary']['total_profit'] for c in app._bulk_results['configurations']], [2, 4, 6])
                self.assertEqual(app.bulk_progress_var.get(), 100)
                snapshots = list(Path(directory).glob('bulk_test_*.json'))
                self.assertEqual(len(snapshots), 1)
                self.assertTrue(snapshots[0].with_suffix('.xlsx').exists())
                self.assertEqual(list(Path(directory).glob('*.sqlite')), [])
                saved = json.loads(snapshots[0].read_text(encoding='utf-8'))
                self.assertEqual(saved['status'], 'completed')
                self.assertNotIn('seed', json.dumps(saved))
                with patch('tkinter.filedialog.askopenfilename', return_value=str(snapshots[0])), \
                     patch.object(app, '_show_bulk_grid_popup') as show_loaded:
                    app.load_bulk_results()
                    show_loaded.assert_called_once()
                self.assertEqual(app._bulk_results, saved)
                app._show_bulk_grid_popup(saved)
                root.update_idletasks()
                popup = next(w for w in root.winfo_children() if isinstance(w, tk.Toplevel))
                tabs = next(w for w in popup.winfo_children() if isinstance(w, ttk.Notebook))
                self.assertEqual(tabs.tab(tabs.tabs()[0], 'text'), 'Confronto laptop')
                self.assertEqual(len(tabs.tabs()), 2)
                popup.destroy()

                app.show_tile_analysis()
                root.update_idletasks()
                analysis = next(w for w in root.winfo_children() if isinstance(w, tk.Toplevel))
                self.assertIn('Analisi', analysis.title())
                analysis.destroy()
                app.tile_record_enabled_var.set(False)
                app.tile_use_history_var.set(True)
                app._update_tile_controls()
                self.assertFalse(app.tile_use_history_var.get())
                self.assertIn('disabled', app.tile_history_check.state())
                app.tile_record_enabled_var.set(True)
                app._update_tile_controls()
                self.assertNotIn('disabled', app.tile_history_check.state())

            root.after_cancel(app.mouse_monitor_after_id)
            position.side_effect = RuntimeError('Position unavailable')
            app.start_mouse_monitoring()
            root.update()
            self.assertIsNone(app.mouse_monitor_after_id)
            self.assertIn('Position unavailable', app.status_text.get('1.0', 'end'))
            position.side_effect = None
            app.start_mouse_monitoring()
            timer = app.mouse_monitor_after_id
            with patch.object(root, 'after_cancel', wraps=root.after_cancel) as cancel:
                app.on_closing()
                cancel.assert_called_once_with(timer)
            self.assertFalse(app.running_event.is_set())
            self.assertEqual(callback_errors, [])

    def test_async_runner(self):
        engine = GameEngine(None, None)
        with patch.object(engine, 'main_game_loop', new_callable=AsyncMock) as loop:
            engine.run_game_async()
            loop.assert_awaited_once()
        with self.assertRaises(RuntimeError):
            asyncio.get_running_loop()
