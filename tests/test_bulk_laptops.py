"""Laptop independence and the financial figures presented to the user."""

from copy import deepcopy
import json
import random
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from bulk_export import laptop_totals, write_bulk_excel, save_bulk_json, load_bulk_json
from bulk_simulation import run_bulk_batch, build_configurations
from test_bulk_simulation import snapshot


class LaptopTest(unittest.TestCase):
    def test_identical_laptops_have_independent_random_streams(self):
        settings = snapshot(target_profit=100, max_loss=90, max_rounds=40,
                            mine_config={d: 8 for d in ("low", "medium", "high")})
        laptops = [deepcopy(settings) for _ in range(3)]
        for index, config in enumerate(laptops, 1):
            config["label"] = f"Laptop {index}"
        before = random.getstate()
        args = (laptops, 100, 314159, lambda *_: None, lambda _: None, lambda: False, lambda: False)
        first = run_bulk_batch(*args)
        second = run_bulk_batch(*args)
        self.assertEqual(random.getstate(), before)
        # Repeatable internal test harness; no RNG controls or identifiers in reports.
        self.assertEqual(first["configurations"], second["configurations"])
        paths = [[r["profit"] for r in c["runs"]] for c in first["configurations"]]
        self.assertNotEqual(paths[0], paths[1])
        self.assertNotEqual(paths[0], paths[2])
        self.assertNotEqual(paths[1], paths[2])
        self.assertNotIn('"seed"', json.dumps(first))
        self.assertEqual(laptops[0]["starting_cash"], settings["starting_cash"])

    def test_gains_minus_losses_equals_net_excluding_interruption(self):
        config = {"runs": [{"profit": 3, "status": "target"}, {"profit": 4, "status": "target"},
                            {"profit": -20, "status": "max_loss"}, {"profit": 999, "status": "interrupted"}]}
        values = laptop_totals(config)
        self.assertEqual(values["completed"], 3)
        self.assertEqual((values["gains"], values["losses"], values["net"]), (7, 20, -13))
        self.assertAlmostEqual(values["average"], -13/3)
        self.assertIsNone(laptop_totals({"runs": []})["net"])

    def test_simple_excel_and_original_color_scale(self):
        settings = [snapshot(label=f"Laptop {i}", target_profit=i) for i in (1, 2, 3)]
        result = run_bulk_batch(settings, 2, 123, lambda *_: None, lambda _: None, lambda: False, lambda: False)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/"report.json"
            save_bulk_json(path, result)
            self.assertEqual(load_bulk_json(path), json.loads(json.dumps(result)))
            write_bulk_excel(path.with_suffix(".xlsx"), result)
            book = load_workbook(path.with_suffix(".xlsx"))
            try:
                sheet = book["Confronto laptop"]
                self.assertEqual(book.sheetnames[0], "Confronto laptop")
                self.assertEqual([sheet.cell(r, 5).value for r in range(6, 9)], [2, 4, 6])
                self.assertEqual(sheet["D6"].value, 0)
                grid = book["Griglia Win Rate"]
                rules = next(iter(grid.conditional_formatting._cf_rules.values()))
                colors = [color.rgb[-6:] for color in rules[0].colorScale.color]
                self.assertEqual(colors, ["F8696B", "FFEB84", "63BE7B"])
                for ws in book:
                    self.assertFalse(any("seed" in str(cell.value).lower() for row in ws for cell in row))
            finally:
                book.close()
            self.assertEqual(list(Path(directory).glob("*.sqlite")), [])

    def test_removed_mode_cannot_be_selected(self):
        with self.assertRaises(ValueError):
            build_configurations(snapshot(mode="fixed"))


if __name__ == "__main__":
    unittest.main()
